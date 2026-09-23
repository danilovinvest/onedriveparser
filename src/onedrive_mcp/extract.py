"""Turn downloaded file bytes into plain text Claude can read."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from pathlib import PurePosixPath

import docx
import openpyxl
from pypdf import PdfReader

TEXT_EXTENSIONS = frozenset(
    {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml",
     ".html", ".htm", ".log", ".py", ".js", ".ts", ".ini", ".toml"}
)


# Office files are zips: a small download can inflate to gigabytes.
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_PDF_PAGES = 500


class UnsupportedFormat(ValueError):
    """Raised when no extractor exists for a file type."""


class ExtractionLimit(ValueError):
    """Raised when a document would inflate beyond safe bounds."""


def check_zip_size(data: bytes, limit: int = MAX_UNCOMPRESSED_BYTES) -> None:
    # ZipExtFile never yields more than the declared file_size, so the sum of
    # declared sizes is a real bound on what the parsers can decompress.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        total = sum(info.file_size for info in archive.infolist())
    if total > limit:
        raise ExtractionLimit(
            f"Archive inflates to {total} bytes, above the {limit} byte limit"
        )


def _pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    count = len(reader.pages)
    pages = (reader.pages[i].extract_text() or "" for i in range(min(count, MAX_PDF_PAGES)))
    text = "\n\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1))
    if count > MAX_PDF_PAGES:
        text += f"\n\n[stopped after {MAX_PDF_PAGES} of {count} pages]"
    return text


def _docx(data: bytes) -> str:
    check_zip_size(data)
    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(lines)


def _cell(value: object) -> str:
    return "" if value is None else str(value)


def _xlsx(data: bytes) -> str:
    check_zip_size(data)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            rows = (
                "\t".join(_cell(v) for v in row)
                for row in sheet.iter_rows(values_only=True)
            )
            sheets.append(f"--- sheet {sheet.title} ---\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        workbook.close()


def _text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


_BY_EXTENSION: dict[str, Callable[[bytes], str]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".xlsx": _xlsx,
    **{ext: _text for ext in TEXT_EXTENSIONS},
}


def extract_text(name: str, data: bytes, mime_type: str | None = None) -> str:
    extractor = _BY_EXTENSION.get(PurePosixPath(name).suffix.lower())
    if extractor is None and mime_type and mime_type.startswith("text/"):
        extractor = _text
    if extractor is None:
        raise UnsupportedFormat(
            f"No text extractor for {name!r} ({mime_type or 'unknown type'}). "
            "Supported: pdf, docx, xlsx and plain-text formats."
        )
    return extractor(data)


def truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + f"\n\n[truncated: {len(text) - max_chars} more characters]"
