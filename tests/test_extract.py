import io
import zipfile

import docx
import openpyxl
import pytest

from onedrive_mcp.extract import (
    ExtractionLimit,
    UnsupportedFormat,
    check_zip_size,
    extract_text,
    truncate,
)


def test_plain_text_by_extension_and_mime() -> None:
    assert extract_text("notes.MD", "héllo".encode()) == "héllo"
    assert extract_text("noext", b"hi", "text/plain") == "hi"


def test_docx_paragraphs_and_tables() -> None:
    document = docx.Document()
    document.add_paragraph("Bonjour")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "a", "b"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("f.docx", buffer.getvalue())
    assert "Bonjour" in text and "a | b" in text


def test_xlsx_sheets() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["item", 12])
    buffer = io.BytesIO()
    workbook.save(buffer)
    text = extract_text("f.xlsx", buffer.getvalue())
    assert "--- sheet Budget ---" in text and "item\t12" in text


def test_unsupported_format() -> None:
    with pytest.raises(UnsupportedFormat):
        extract_text("photo.jpg", b"\xff\xd8", "image/jpeg")


def test_truncate() -> None:
    assert truncate("abc", 5) == "abc"
    assert truncate("abcdef", 2).startswith("ab\n\n[truncated: 4 more")


def test_zip_bomb_is_rejected_before_parsing() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * 2_000_000)
    data = buffer.getvalue()
    assert len(data) < 10_000  # tiny download, large inflation
    with pytest.raises(ExtractionLimit):
        check_zip_size(data, limit=1_000_000)
    check_zip_size(data, limit=3_000_000)


def test_pdf_page_cap(monkeypatch: pytest.MonkeyPatch) -> None:
    from pypdf import PdfWriter

    from onedrive_mcp import extract

    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    monkeypatch.setattr(extract, "MAX_PDF_PAGES", 2)
    text = extract_text("f.pdf", buffer.getvalue())
    assert "--- page 2 ---" in text and "--- page 3 ---" not in text
    assert text.endswith("[stopped after 2 of 3 pages]")
